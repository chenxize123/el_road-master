from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "project_promo_assets"
OUT_DOCX = OUT_DIR / "智学路_项目宣传与介绍文档.docx"

FONT_REG = Path("C:/Windows/Fonts/msyh.ttc")
FONT_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")

BLUE = "1F5EFF"
NAVY = "17324D"
CYAN = "24B8DB"
GREEN = "29A66A"
GOLD = "F2A900"
RED = "E85D75"
MUTED = "6B7280"
LIGHT_BLUE = "EAF2FF"
LIGHT_CYAN = "EAFBFF"
LIGHT_GRAY = "F6F8FB"
WHITE = "FFFFFF"


def font(size, bold=False):
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REG
    return ImageFont.truetype(str(path), size)


def rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


def make_rounded_rect(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def wrap_text(draw, text, fnt, max_width):
    lines = []
    line = ""
    for ch in text:
        trial = line + ch
        if draw.textlength(trial, font=fnt) <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = ch
    if line:
        lines.append(line)
    return lines


def draw_text_block(draw, text, xy, fnt, fill, max_width, line_gap=8):
    x, y = xy
    for line in wrap_text(draw, text, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def create_loop_image(path):
    img = Image.new("RGB", (1600, 900), rgb("F8FBFF"))
    d = ImageDraw.Draw(img)
    title_f = font(54, True)
    sub_f = font(24)
    node_f = font(30, True)
    small_f = font(21)

    d.text((80, 62), "AI 智慧助教闭环", font=title_f, fill=rgb(NAVY))
    d.text((82, 132), "围绕课程建设、班级教学、学习互动、测评反馈形成可运营的数据闭环", font=sub_f, fill=rgb(MUTED))

    center = (800, 480)
    nodes = [
        ("课程建设", "课程信息、数字人视频、知识图谱", 800, 210, BLUE),
        ("班级教学", "功能开关、学生管理、教学资源", 1145, 400, CYAN),
        ("学习互动", "讨论区、实验提交、代码问答", 1020, 705, GREEN),
        ("测评考试", "题库、试卷、考试、成绩结果", 580, 705, GOLD),
        ("智能分析", "AI 总结、讨论评分、实验分析", 455, 400, RED),
    ]

    # Flow arrows
    for i, (_, _, x1, y1, c1) in enumerate(nodes):
        _, _, x2, y2, _ = nodes[(i + 1) % len(nodes)]
        d.line((x1, y1, x2, y2), fill=rgb("C9D6EA"), width=7)
        # Simple arrow head
        hx = x1 * 0.22 + x2 * 0.78
        hy = y1 * 0.22 + y2 * 0.78
        d.ellipse((hx - 10, hy - 10, hx + 10, hy + 10), fill=rgb(c1))

    d.ellipse((center[0] - 170, center[1] - 170, center[0] + 170, center[1] + 170), fill=rgb("FFFFFF"), outline=rgb("D6E3F7"), width=4)
    d.text((center[0] - 104, center[1] - 42), "智学路", font=font(58, True), fill=rgb(BLUE))
    d.text((center[0] - 122, center[1] + 30), "大模型智慧助教系统", font=font(27, True), fill=rgb(NAVY))

    for title, desc, x, y, color in nodes:
        make_rounded_rect(d, (x - 190, y - 72, x + 190, y + 72), 28, rgb("FFFFFF"), rgb(color), 4)
        d.ellipse((x - 174, y - 18, x - 138, y + 18), fill=rgb(color))
        d.text((x - 120, y - 42), title, font=node_f, fill=rgb(NAVY))
        d.text((x - 120, y + 6), desc, font=small_f, fill=rgb(MUTED))

    img.save(path, quality=95)


def create_journey_image(path):
    img = Image.new("RGB", (1600, 900), rgb("FFFFFF"))
    d = ImageDraw.Draw(img)
    title_f = font(52, True)
    role_f = font(33, True)
    step_f = font(24, True)
    desc_f = font(20)

    d.text((80, 62), "三类角色协同的教学运营旅程", font=title_f, fill=rgb(NAVY))
    d.text((82, 130), "从建课、授课、学习到评价沉淀，减少教师重复劳动，提升学生学习反馈质量", font=font(24), fill=rgb(MUTED))

    roles = [
        ("管理员", BLUE, ["用户与权限", "课程/班级监管", "平台运营"]),
        ("教师", GREEN, ["创建课程", "发布资源与考试", "AI 辅助评价"]),
        ("学生", GOLD, ["进入班级", "学习与讨论", "考试/实验提交"]),
    ]
    y_positions = [255, 455, 655]
    x0 = 115
    step_w = 330
    gap = 48

    for (role, color, steps), y in zip(roles, y_positions):
        make_rounded_rect(d, (x0, y - 62, x0 + 210, y + 62), 26, rgb(color), None)
        d.text((x0 + 48, y - 22), role, font=role_f, fill=rgb(WHITE))
        for idx, step in enumerate(steps):
            x = x0 + 260 + idx * (step_w + gap)
            make_rounded_rect(d, (x, y - 62, x + step_w, y + 62), 24, rgb(LIGHT_GRAY), rgb("DDE5F1"), 2)
            d.text((x + 28, y - 34), f"{idx + 1}. {step}", font=step_f, fill=rgb(NAVY))
            desc = [
                "按角色开放能力边界",
                "沉淀课堂与学习过程数据",
                "可持续改进教学效果",
            ][idx]
            d.text((x + 30, y + 10), desc, font=desc_f, fill=rgb(MUTED))
            if idx < 2:
                x_arrow = x + step_w + 12
                d.line((x_arrow, y, x_arrow + 26, y), fill=rgb(color), width=5)
                d.polygon([(x_arrow + 26, y), (x_arrow + 12, y - 9), (x_arrow + 12, y + 9)], fill=rgb(color))

    img.save(path, quality=95)


def create_module_image(path):
    img = Image.new("RGB", (1600, 900), rgb("F8FAFC"))
    d = ImageDraw.Draw(img)
    d.text((80, 60), "核心模块一览", font=font(52, True), fill=rgb(NAVY))
    d.text((82, 128), "以课程为中心，以班级为教学现场，以 AI 能力增强内容生产、互动分析和评价反馈", font=font(24), fill=rgb(MUTED))

    cards = [
        ("课程管理", "建课、课程资料、封面、课程班级关联", BLUE),
        ("数字人视频", "教学视频上传、播放、资源列表管理", CYAN),
        ("知识图谱", "知识点层级、关系展示、节点详情", GREEN),
        ("题库与试卷", "题库管理、组卷、题目即入库", GOLD),
        ("考试测评", "考试发布、时间控制、成绩结果", RED),
        ("实验分析", "实验发布、代码提交、AI 评分与问答", "7C3AED"),
        ("讨论区", "主题讨论、AI 总结、过程性评分", "0EA5E9"),
        ("权限控制", "教师、学生、管理员、访客角色隔离", "64748B"),
    ]

    x_start, y_start = 92, 230
    card_w, card_h = 335, 190
    gap_x, gap_y = 42, 44
    for idx, (title, desc, color) in enumerate(cards):
        row, col = divmod(idx, 4)
        x = x_start + col * (card_w + gap_x)
        y = y_start + row * (card_h + gap_y)
        make_rounded_rect(d, (x, y, x + card_w, y + card_h), 24, rgb("FFFFFF"), rgb("E2E8F0"), 2)
        d.rectangle((x, y, x + card_w, y + 14), fill=rgb(color))
        d.ellipse((x + 26, y + 38, x + 72, y + 84), fill=rgb(color))
        d.text((x + 92, y + 38), title, font=font(29, True), fill=rgb(NAVY))
        draw_text_block(d, desc, (x + 30, y + 108), font(22), rgb(MUTED), card_w - 60, 8)

    img.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2F0", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=None, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", size=11, bold=False, color="1F2937", after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=17 if level == 1 else 13, bold=True, color=BLUE if level == 1 else NAVY)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.16
    r = p.add_run(text)
    set_font(r, size=10.5, color="374151")
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "BFD2F6", "8")
    set_cell_margins(cell, 160, 180, 160, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=12, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_font(r2, size=10.5, color="374151")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(hdr[i], "E8EEF5")
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, 10.2, True, NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_font(r, 9.5, False, "374151")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智学路大模型智慧助教系统 | 项目宣传与介绍")
    set_font(r, 9, False, MUTED)


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = styles[name]
        st.font.name = "微软雅黑"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    # Real list style tuning.
    bullet = styles["List Bullet"]
    bullet.font.name = "微软雅黑"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    bullet.paragraph_format.left_indent = Inches(0.26)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(4)


def add_cover(doc):
    banner = ROOT / "src" / "assets" / "images" / "dashboard-banner.png"
    if banner.exists():
        doc.add_picture(str(banner), width=Inches(6.7))
    add_para(doc, "智学路", size=31, bold=True, color=NAVY, before=22, after=0)
    add_para(doc, "大模型智慧助教系统", size=22, bold=True, color=BLUE, after=10)
    add_para(
        doc,
        "面向高校与职业教育场景的一体化课程建设、班级教学、AI 互动分析与测评管理平台",
        size=13,
        color="374151",
        after=18,
    )
    add_callout(
        doc,
        "给意向方的核心信息",
        "智学路把课程资源、知识图谱、题库试卷、考试、实验、讨论与 AI 分析整合到同一教学工作台，帮助学校把“内容建设、过程管理、学习反馈、质量评价”从分散工具转为可持续运营的闭环。",
        LIGHT_CYAN,
    )
    rows = [
        ("适用对象", "高校学院、职业院校、培训机构、教学改革项目团队"),
        ("核心价值", "降低教师重复建设成本，增强学生学习反馈，沉淀课程与班级过程数据"),
        ("部署形态", "前后端分离架构，支持按学校/课程/班级进行功能配置与权限隔离"),
        ("当前版本", "已具备课程、班级、题库、试卷、考试、实验、讨论、知识图谱与 AI 分析能力"),
    ]
    add_table(doc, ["项目项", "说明"], rows, [1.35, 5.15])
    doc.add_page_break()


def add_body(doc):
    add_heading(doc, "1. 项目定位", 1)
    add_para(
        doc,
        "智学路是一套面向现代教学场景的智慧助教平台。它不是单一的考试系统，也不是单纯的资源库，而是以课程为核心，把教师备课、班级管理、学生学习、过程互动、实验训练、考试测评和 AI 辅助分析连接起来。",
        after=8,
    )
    add_bullet(doc, "对学校：形成统一的课程数字化建设与教学质量观察入口。")
    add_bullet(doc, "对教师：把课程资源、题库、试卷、实验和讨论集中管理，减少重复操作。")
    add_bullet(doc, "对学生：在班级空间内完成学习、讨论、考试、实验提交与代码问答。")
    add_bullet(doc, "对管理者：通过角色、权限和班级功能开关，控制不同课程的开放能力。")

    loop_path = ASSET_DIR / "ai_teaching_loop.png"
    doc.add_picture(str(loop_path), width=Inches(6.5))
    cap = add_para(doc, "图 1：围绕课程与班级形成的 AI 智慧助教闭环", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    add_heading(doc, "2. 产品能力", 1)
    module_path = ASSET_DIR / "module_overview.png"
    doc.add_picture(str(module_path), width=Inches(6.5))
    add_para(doc, "图 2：核心模块与能力范围", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    add_table(
        doc,
        ["模块", "面向意向方的价值", "已体现能力"],
        [
            ("课程管理", "承载课程基础数据与教学资源，形成可复用课程资产。", "课程创建、课程详情、课程列表、课程封面与分享入口"),
            ("班级管理", "把课程落到真实教学组织中，支持分班级运营与功能开放。", "班级列表、学生管理、班级详情、功能控制器"),
            ("知识图谱", "将章节、知识点和前置关系可视化，帮助学生理解知识结构。", "图谱展示、节点详情、关系浏览"),
            ("数字人视频", "集中管理教学视频资源，支撑异步学习与课程内容沉淀。", "视频上传、列表、播放、异常提示"),
            ("题库/试卷/考试", "让教师从题目沉淀到组卷、发布考试、查看结果形成闭环。", "题库管理、题目管理、组卷、考试时间与班级关联"),
            ("实验与代码问答", "适合程序设计、工程实践类课程，支持提交、分析、问答追踪。", "实验创建、代码提交、AI 分析、得分与问答状态"),
            ("讨论区", "把课堂讨论转为可分析的过程性学习证据。", "讨论主题、发表内容、AI 分析/总结、讨论评分、Excel 导出"),
        ],
        [1.15, 2.55, 2.8],
    )

    add_heading(doc, "3. 典型使用场景", 1)
    journey_path = ASSET_DIR / "role_journey.png"
    doc.add_picture(str(journey_path), width=Inches(6.5))
    add_para(doc, "图 3：管理员、教师、学生三类角色的协同旅程", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    add_heading(doc, "面向课程负责人", 2)
    add_para(doc, "课程负责人可以快速创建课程、维护课程资源、构建知识图谱、沉淀题库，并通过课程控制台把题库、试卷、考试和课程资料串联起来。")
    add_heading(doc, "面向授课教师", 2)
    add_para(doc, "授课教师可以围绕班级组织教学活动：打开或关闭知识图谱、数字人视频、考试、实验、课程赋能课题、讨论区等功能，并对学生学习过程进行管理。")
    add_heading(doc, "面向学生", 2)
    add_para(doc, "学生在班级空间中进入课程资源、参与讨论、查看知识图谱、完成考试与实验提交。实验模块支持代码问答，帮助学生在提交后进一步理解问题与改进方向。")

    doc.add_page_break()

    add_heading(doc, "4. AI 赋能亮点", 1)
    add_callout(
        doc,
        "AI 不是孤立入口，而是嵌入教学流程",
        "平台将 AI 能力放在讨论、实验、题目生成和学习反馈等实际场景中，避免“只聊天、不沉淀”的工具化使用，让 AI 结果回到课程、班级、题库和评价数据中。",
        LIGHT_BLUE,
    )
    add_bullet(doc, "AI 出题：支持按知识点生成填空、判断、单选、多选等题型，并服务题库沉淀。")
    add_bullet(doc, "讨论分析：支持对讨论内容进行 AI 分析与总结，并保留到讨论详情中。")
    add_bullet(doc, "过程评分：讨论区具备参与度、内容质量、认知深度等维度评分数据，并支持导出。")
    add_bullet(doc, "实验分析：实验提交后可结合 AI 分析得分、原因说明与代码问答，帮助教师识别掌握情况。")
    add_bullet(doc, "知识结构化：知识图谱帮助将碎片化课程内容组织成可浏览、可解释的学习路径。")

    add_heading(doc, "5. 技术与集成基础", 1)
    add_para(doc, "当前前端采用 Vue 3、Vite 与 Ant Design Vue 构建，接口层通过统一 API 服务处理认证、超时、Bearer Token 与大整数 JSON 精度问题。业务接口覆盖课程、班级、用户、题库、试卷、考试、实验、讨论、文件与权限等模块。")
    add_table(
        doc,
        ["技术维度", "说明"],
        [
            ("前端架构", "Vue 3 + Vite + Ant Design Vue，组件化组织课程、班级、考试、讨论等页面。"),
            ("权限模型", "路由与功能入口按 admin、teacher、student、visitor 等角色区分。"),
            ("数据接口", "统一 API 服务封装 GET/POST/PUT/DELETE，支持 Token 注入和超时处理。"),
            ("可配置教学", "班级控制器支持按班级开放知识图谱、视频、考试、实验、讨论等能力。"),
            ("可扩展模块", "课程和班级均采用标签页模块化结构，便于后续扩展新教学能力。"),
        ],
        [1.4, 5.1],
    )

    add_heading(doc, "6. 合作落地建议", 1)
    add_table(
        doc,
        ["阶段", "建议目标", "可交付成果"],
        [
            ("试点准备", "选择 1-3 门适合数字化改革的课程，梳理角色、班级、资源与题库。", "试点课程清单、功能开关策略、账号角色方案"),
            ("课程建设", "完成课程基本信息、视频资源、知识图谱、题库与试卷建设。", "课程空间、知识图谱、题库资产、试卷样例"),
            ("教学运行", "在真实班级中运行讨论、考试、实验提交与 AI 分析流程。", "学生学习记录、讨论评分、实验提交与考试数据"),
            ("复盘推广", "对试点课程的教学效果、教师工作量与学生反馈进行复盘。", "试点报告、推广清单、二期功能优化建议"),
        ],
        [1.0, 2.75, 2.75],
    )

    add_heading(doc, "7. 项目优势总结", 1)
    add_bullet(doc, "一体化：课程、班级、考试、实验、讨论和知识图谱在同一平台内联动。")
    add_bullet(doc, "教学真实可用：覆盖教师日常建课、授课、出题、组卷、发布考试、查看提交等高频操作。")
    add_bullet(doc, "AI 场景明确：AI 能力服务于出题、讨论总结、实验分析、代码问答等具体流程。")
    add_bullet(doc, "过程数据可沉淀：讨论评分、实验分析、考试结果等数据可作为教学质量改进依据。")
    add_bullet(doc, "扩展空间清晰：模块化前端与统一接口层便于继续接入更多教学智能体和数据看板。")

    add_callout(
        doc,
        "推荐沟通口径",
        "智学路适合用于高校智慧课程建设、人工智能赋能教学改革、程序设计/工程实践类课程过程评价、课程资源与题库资产沉淀等场景。它的价值不只在“有 AI”，而在于把 AI 放进真实教学工作流，并形成可复用、可观察、可改进的教学闭环。",
        LIGHT_CYAN,
    )


def main():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_loop_image(ASSET_DIR / "ai_teaching_loop.png")
    create_journey_image(ASSET_DIR / "role_journey.png")
    create_module_image(ASSET_DIR / "module_overview.png")

    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    add_body(doc)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
